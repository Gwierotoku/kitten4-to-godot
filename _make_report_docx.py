# -*- coding: utf-8 -*-
"""把技术路线汇报的 Markdown 转成 Word 文档（结构化排版，非逐字转换）。"""
import re
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SRC = r"D:\DS-projects\kitten4转Godot项目\技术路线汇报.md"
DST = r"D:\DS-projects\kitten4转Godot项目\技术路线汇报.docx"

CN_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"


def set_font(run, name=CN_FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_inline(par, text, base_size=10.5, base_bold=False):
    """处理 **粗体** 与 `代码` 两种行内标记。"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            set_font(r, CN_FONT, base_size, True)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            set_font(r, MONO_FONT, base_size - 0.5, False,
                     RGBColor(0xB0, 0x30, 0x60))
        else:
            r = par.add_run(part)
            set_font(r, CN_FONT, base_size, base_bold)


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("\n".join(lines))
    set_font(r, MONO_FONT, 9, False, RGBColor(0x22, 0x44, 0x66))


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            cell = cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline(p, txt, base_size=9, base_bold=(i == 0))
    doc.add_paragraph()


def main():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = CN_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        # 代码块
        if s.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            add_code_block(doc, buf)
            i += 1
            continue

        # 表格
        if s.startswith("|") and i + 1 < len(lines) and re.match(
                r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = []
            hdr = [c.strip() for c in s.strip("|").split("|")]
            rows.append(hdr)
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1))
            text = m.group(2)
            if lvl == 1:
                p = doc.add_heading("", level=0)
                add_inline(p, text, base_size=20, base_bold=True)
            else:
                p = doc.add_heading("", level=lvl - 1)
                size = {2: 15, 3: 13, 4: 11.5}.get(lvl, 11)
                add_inline(p, text, base_size=size, base_bold=True)
            i += 1
            continue

        # 引用
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, s.lstrip("> ").strip(), base_size=10)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # 分隔线
        if re.match(r"^-{3,}$", s):
            p = doc.add_paragraph()
            r = p.add_run("─" * 40)
            set_font(r, CN_FONT, 8, False, RGBColor(0xBB, 0xBB, 0xBB))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 空行
        if s == "":
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            ordered = m.group(2).endswith(".")
            style = "List Number" if ordered else "List Bullet"
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph()
                txt = ("• " if not ordered else "") + m.group(3)
                p.paragraph_format.left_indent = Cm(0.75 + indent * 0.6)
                add_inline(p, txt)
                i += 1
                continue
            p.paragraph_format.left_indent = Cm(0.75 + indent * 0.6)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(3))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        add_inline(p, s)
        i += 1

    doc.save(DST)
    print("saved:", DST)
    print("size:", os.path.getsize(DST), "bytes")


if __name__ == "__main__":
    main()
